"""
Giao diện thống kê bảng xếp hạng dự đoán bóng đá.

Cài đặt: pip install streamlit openpyxl pandas
Chạy:    streamlit run dashboard.py
"""

import json
import re

import pandas as pd
import streamlit as st
from openpyxl import load_workbook

POLL_FILE   = "poll_results_1.xlsx"
SCORES_FILE = "scores.json"

st.set_page_config(
    page_title="Dự đoán bóng đá",
    page_icon="⚽",
    layout="wide",
)

# ─────────────────────────────────────────────────────────────
# Load dữ liệu
# ─────────────────────────────────────────────────────────────

@st.cache_data
def load_all():
    wb = load_workbook(POLL_FILE, data_only=True)
    ws = wb["Kết quả vote"]
    rows        = list(ws.iter_rows(values_only=True))
    headers     = rows[1]
    match_names = [h for h in headers[2:] if h]

    with open(SCORES_FILE, encoding="utf-8") as f:
        raw = json.load(f)
    scores = {k: tuple(v) for k, v in raw.items() if v is not None}

    members = []
    for row in rows[2:]:
        if not row[1]:
            continue
        if row[0] == "Tổng":  # hàng tổng cộng cuối sheet
            continue
        name = str(row[1]).split("\n")[0].strip()
        votes = {
            match: (row[col + 2] if (col + 2) < len(row) else None)
            for col, match in enumerate(match_names)
        }
        members.append({"name": name, "votes": votes})

    return match_names, scores, members


# ─────────────────────────────────────────────────────────────
# Logic chấm điểm
# ─────────────────────────────────────────────────────────────

def _resolve_team(team: str, home: str, away: str) -> str | None:
    """Khớp tên đội — exact trước, rồi prefix (xử lý 'Netherland', 'Bosnia', ...)."""
    if team == home: return home
    if team == away: return away
    # prefix match: "Netherland" → "Netherlands", "Bosnia" → "Bosnia-Herzegovina"
    for candidate in (home, away):
        if candidate.startswith(team) or team.startswith(candidate):
            return candidate
    return None


def grade(pred, match, scores) -> bool | None:
    """True = đúng, False = sai, None = không vote / không parse được."""
    if pred is None or str(pred).strip() == "":
        return None
    pred = str(pred).strip()
    m = re.match(r"(.+?)\s*([+-]\d+(?:\.\d+)?)$", pred)
    if not m:
        return None
    team, handicap = m.group(1).strip(), float(m.group(2))
    home, away = match.split(" vs ")
    hs, as_ = scores[match]
    resolved = _resolve_team(team, home, away)
    if resolved == home:
        diff = hs + handicap - as_
    elif resolved == away:
        diff = as_ + handicap - hs
    else:
        return None
    return diff > 0


def build_ranking(members: list, scores: dict) -> pd.DataFrame:
    rows = []
    for m in members:
        scored  = {k: v for k, v in m["votes"].items() if k in scores}
        total   = len(scored)
        correct = sum(1 for match, pred in scored.items() if grade(pred, match, scores) is True)
        rows.append({
            "Tên":       m["name"],
            "Đúng":      correct,
            "Sai":       total - correct,
            "Tỷ lệ (%)": round(correct / total * 100, 1) if total else 0,
        })
    df = (
        pd.DataFrame(rows)
        .sort_values(["Đúng", "Sai", "Tên"], ascending=[False, True, True])
        .reset_index(drop=True)
    )
    df.index     += 1
    df.index.name = "Hạng"
    return df


def build_detail(member: dict, scores: dict) -> pd.DataFrame:
    rows = []
    for match, pred in member["votes"].items():
        if match not in scores:
            continue
        home, away = match.split(" vs ")
        hs, as_    = scores[match]
        result     = grade(pred, match, scores)
        rows.append({
            "Trận":    match,
            "Tỉ số":  f"{hs} - {as_}",
            "Dự đoán": str(pred).strip() if pred else "—",
            "Kết quả": "✅ Đúng" if result is True else ("❌ Sai" if result is False else "⬜ Không vote"),
        })
    return pd.DataFrame(rows)


# ─────────────────────────────────────────────────────────────
# Xuất HTML tĩnh — self-contained, không cần server
# ─────────────────────────────────────────────────────────────

def export_html(
    match_names: list,
    scores: dict,
    members: list,
    ranking_df: pd.DataFrame,
    member_map: dict,
    output_path: str | None = None,
) -> bytes:
    """
    Tạo file HTML tự chứa (không cần server) có thể gửi cho người khác.
    Trả về bytes; nếu output_path được cung cấp thì cũng ghi ra file.
    """

    # Chuẩn bị dữ liệu cho JS
    ranking_list = [
        {
            "rank":     int(ranking_df.index[i]),
            "name":     ranking_df.iloc[i]["Tên"],
            "correct":  int(ranking_df.iloc[i]["Đúng"]),
            "wrong":    int(ranking_df.iloc[i]["Sai"]),
            "accuracy": float(ranking_df.iloc[i]["Tỷ lệ (%)"]),
        }
        for i in range(len(ranking_df))
    ]

    details_map: dict[str, list] = {}
    for name, member in member_map.items():
        rows = []
        for match, pred in member["votes"].items():
            if match not in scores:
                continue
            home, away = match.split(" vs ")
            hs, as_ = scores[match]
            result = grade(pred, match, scores)
            rows.append({
                "match":   match,
                "score":   f"{hs} - {as_}",
                "pred":    str(pred).strip() if pred else "—",
                "result":  "correct" if result is True else ("wrong" if result is False else "novote"),
            })
        details_map[name] = rows

    data_json = json.dumps(
        {"ranking": ranking_list, "details": details_map},
        ensure_ascii=False,
    )

    avg_acc   = ranking_df["Tỷ lệ (%)"].mean()
    top_name  = ranking_df.iloc[0]["Tên"] if not ranking_df.empty else "—"
    generated = __import__("datetime").datetime.now().strftime("%d/%m/%Y %H:%M")

    html = f"""<!DOCTYPE html>
<html lang="vi">
<head>
<meta charset="UTF-8"/>
<meta name="viewport" content="width=device-width, initial-scale=1.0"/>
<title>⚽ Bảng xếp hạng dự đoán bóng đá</title>
<style>
  *, *::before, *::after {{ box-sizing: border-box; margin: 0; padding: 0; }}
  body {{ font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", sans-serif;
          background: #f0f2f6; color: #1a1a2e; min-height: 100vh; }}

  header {{ background: linear-gradient(135deg, #1F4E79, #2e86de);
            color: #fff; padding: 20px 32px; display: flex;
            align-items: center; gap: 12px; box-shadow: 0 2px 8px rgba(0,0,0,.2); }}
  header h1 {{ font-size: 1.5rem; font-weight: 700; }}
  header small {{ font-size: .8rem; opacity: .75; margin-left: auto; }}

  .metrics {{ display: grid; grid-template-columns: repeat(4, 1fr);
              gap: 14px; padding: 20px 24px 10px; }}
  .metric {{ background: #fff; border-radius: 10px; padding: 16px 20px;
             box-shadow: 0 1px 4px rgba(0,0,0,.08); border-left: 4px solid #2e86de; }}
  .metric .label {{ font-size: .78rem; color: #666; margin-bottom: 4px; }}
  .metric .value {{ font-size: 1.5rem; font-weight: 700; color: #1F4E79; }}

  .main {{ display: grid; grid-template-columns: 420px 1fr;
           gap: 16px; padding: 10px 24px 24px; }}

  .card {{ background: #fff; border-radius: 12px;
           box-shadow: 0 1px 4px rgba(0,0,0,.08); overflow: hidden; }}
  .card-header {{ background: #1F4E79; color: #fff; padding: 12px 18px;
                  font-weight: 600; font-size: .95rem; }}

  /* Ranking table */
  #ranking-table {{ width: 100%; border-collapse: collapse; font-size: .87rem; }}
  #ranking-table thead th {{ background: #EBF3FB; color: #1F4E79; font-weight: 600;
                              padding: 10px 12px; text-align: left; position: sticky; top: 0; }}
  #ranking-table tbody tr {{ border-bottom: 1px solid #f0f2f6; cursor: pointer;
                              transition: background .15s; }}
  #ranking-table tbody tr:hover {{ background: #EBF3FB; }}
  #ranking-table tbody tr.active {{ background: #d0e8ff; }}
  #ranking-table td {{ padding: 9px 12px; }}
  td.rank {{ font-weight: 700; color: #1F4E79; width: 42px; text-align: center; }}
  td.name {{ font-weight: 500; }}
  td.num  {{ text-align: center; }}

  .bar-wrap {{ background: #f0f2f6; border-radius: 99px; height: 8px;
               width: 80px; overflow: hidden; display: inline-block; vertical-align: middle; }}
  .bar {{ height: 100%; background: #2e86de; border-radius: 99px; }}
  .bar-label {{ margin-left: 6px; font-size: .82rem; color: #444; }}

  .table-scroll {{ max-height: 560px; overflow-y: auto; }}

  /* Detail panel */
  #detail-panel .empty-msg {{ padding: 48px; text-align: center;
                               color: #999; font-size: 1rem; }}
  #detail-meta {{ padding: 16px 18px; border-bottom: 1px solid #f0f2f6; }}
  #detail-meta h2 {{ font-size: 1.15rem; font-weight: 700; margin-bottom: 10px; }}
  .mini-metrics {{ display: flex; gap: 16px; flex-wrap: wrap; }}
  .mini-metric {{ background: #f0f2f6; border-radius: 8px; padding: 10px 16px;
                  min-width: 80px; text-align: center; }}
  .mini-metric .ml {{ font-size: .72rem; color: #666; }}
  .mini-metric .mv {{ font-size: 1.2rem; font-weight: 700; color: #1F4E79; }}

  .tabs {{ display: flex; gap: 0; border-bottom: 2px solid #f0f2f6;
           padding: 0 16px; margin-top: 4px; }}
  .tab {{ padding: 9px 16px; cursor: pointer; font-size: .84rem; border: none;
          background: none; color: #666; border-bottom: 3px solid transparent;
          margin-bottom: -2px; font-weight: 500; transition: color .15s; }}
  .tab.active {{ color: #1F4E79; border-bottom-color: #2e86de; font-weight: 700; }}
  .tab:hover {{ color: #1F4E79; }}

  #detail-table-wrap {{ max-height: 460px; overflow-y: auto; }}
  #detail-table {{ width: 100%; border-collapse: collapse; font-size: .86rem; }}
  #detail-table thead th {{ background: #EBF3FB; color: #1F4E79; font-weight: 600;
                             padding: 10px 12px; text-align: left; position: sticky; top: 0; }}
  #detail-table tbody tr {{ border-bottom: 1px solid #f0f2f6; }}
  #detail-table td {{ padding: 9px 12px; }}
  .result-correct {{ background: #d4edda; color: #155724; font-weight: 600;
                     border-radius: 4px; padding: 2px 8px; white-space: nowrap; }}
  .result-wrong   {{ background: #f8d7da; color: #721c24; font-weight: 600;
                     border-radius: 4px; padding: 2px 8px; white-space: nowrap; }}
  .result-novote  {{ background: #f8f9fa; color: #888;
                     border-radius: 4px; padding: 2px 8px; white-space: nowrap; }}

  .medal {{ display: inline-block; width: 24px; height: 24px; border-radius: 50%;
            text-align: center; line-height: 24px; font-size: .78rem; font-weight: 700;
            margin-right: 4px; }}
  .medal-1 {{ background: #FFD700; color: #7a5c00; }}
  .medal-2 {{ background: #C0C0C0; color: #444; }}
  .medal-3 {{ background: #CD7F32; color: #fff; }}

  @media (max-width: 800px) {{
    .main {{ grid-template-columns: 1fr; }}
    .metrics {{ grid-template-columns: 1fr 1fr; }}
  }}
</style>
</head>
<body>

<header>
  <span style="font-size:1.8rem">⚽</span>
  <h1>Bảng xếp hạng dự đoán bóng đá</h1>
  <small>Xuất lúc {generated}</small>
</header>

<div class="metrics">
  <div class="metric">
    <div class="label">Số trận có kết quả</div>
    <div class="value">{len(scores)}</div>
  </div>
  <div class="metric">
    <div class="label">Số người tham gia</div>
    <div class="value">{len(members)}</div>
  </div>
  <div class="metric">
    <div class="label">Tỷ lệ đúng trung bình</div>
    <div class="value">{avg_acc:.1f}%</div>
  </div>
  <div class="metric">
    <div class="label">Người đứng đầu</div>
    <div class="value" style="font-size:1.1rem">🥇 {top_name}</div>
  </div>
</div>

<div class="main">
  <!-- Ranking -->
  <div class="card">
    <div class="card-header">🏆 Bảng xếp hạng</div>
    <div class="table-scroll">
      <table id="ranking-table">
        <thead>
          <tr>
            <th style="text-align:center">Hạng</th>
            <th>Tên</th>
            <th style="text-align:center">Đúng</th>
            <th style="text-align:center">Sai</th>
            <th>Tỷ lệ</th>
          </tr>
        </thead>
        <tbody id="ranking-body"></tbody>
      </table>
    </div>
  </div>

  <!-- Detail -->
  <div class="card" id="detail-panel">
    <div class="card-header">📋 Chi tiết dự đoán</div>
    <div id="detail-content">
      <div class="empty-msg">👈 Chọn một người trong bảng xếp hạng để xem chi tiết</div>
    </div>
  </div>
</div>

<script>
const DATA = {data_json};

const MEDAL = {{ 1:"medal-1", 2:"medal-2", 3:"medal-3" }};

function medalHtml(rank) {{
  if (rank <= 3) return `<span class="medal ${{MEDAL[rank]}}">${{rank}}</span>`;
  return `<span style="padding:0 4px">${{rank}}</span>`;
}}

// Render ranking table
const tbody = document.getElementById("ranking-body");
DATA.ranking.forEach((r, i) => {{
  const tr = document.createElement("tr");
  tr.dataset.name = r.name;
  const barW = Math.round(r.accuracy);
  tr.innerHTML = `
    <td class="rank">${{medalHtml(r.rank)}}</td>
    <td class="name">${{r.name}}</td>
    <td class="num" style="color:#155724;font-weight:600">${{r.correct}}</td>
    <td class="num" style="color:#721c24">${{r.wrong}}</td>
    <td>
      <div class="bar-wrap"><div class="bar" style="width:${{barW}}%"></div></div>
      <span class="bar-label">${{r.accuracy}}%</span>
    </td>`;
  tr.addEventListener("click", () => showDetail(r.name, r.rank, r.correct, r.wrong, r.accuracy));
  tbody.appendChild(tr);
}});

// Detail panel
let currentTab = "all";

function resultHtml(r) {{
  if (r === "correct") return '<span class="result-correct">✅ Đúng</span>';
  if (r === "wrong")   return '<span class="result-wrong">❌ Sai</span>';
  return '<span class="result-novote">⬜ Không vote</span>';
}}

function renderTable(rows) {{
  if (!rows.length) return '<p style="padding:24px;color:#999;text-align:center">Không có dữ liệu</p>';
  return `<table id="detail-table">
    <thead><tr>
      <th>Trận</th><th style="text-align:center">Tỉ số</th>
      <th>Dự đoán</th><th style="text-align:center">Kết quả</th>
    </tr></thead>
    <tbody>
    ${{rows.map(r => `<tr>
      <td>${{r.match}}</td>
      <td style="text-align:center;font-weight:600">${{r.score}}</td>
      <td>${{r.pred}}</td>
      <td style="text-align:center">${{resultHtml(r.result)}}</td>
    </tr>`).join("")}}
    </tbody></table>`;
}}

function showDetail(name, rank, correct, wrong, accuracy) {{
  // Highlight row
  document.querySelectorAll("#ranking-body tr").forEach(tr => tr.classList.remove("active"));
  document.querySelectorAll("#ranking-body tr").forEach(tr => {{
    if (tr.dataset.name === name) tr.classList.add("active");
  }});

  const rows = DATA.details[name] || [];
  const counts = {{
    all:     rows.length,
    correct: rows.filter(r => r.result === "correct").length,
    wrong:   rows.filter(r => r.result === "wrong").length,
    novote:  rows.filter(r => r.result === "novote").length,
  }};

  currentTab = "all";

  document.getElementById("detail-content").innerHTML = `
    <div id="detail-meta">
      <h2>${{name}}</h2>
      <div class="mini-metrics">
        <div class="mini-metric"><div class="ml">Hạng</div><div class="mv">#${{rank}}</div></div>
        <div class="mini-metric"><div class="ml">Đúng</div><div class="mv" style="color:#155724">${{correct}}</div></div>
        <div class="mini-metric"><div class="ml">Sai</div><div class="mv" style="color:#721c24">${{wrong}}</div></div>
        <div class="mini-metric"><div class="ml">Tỷ lệ</div><div class="mv">${{accuracy}}%</div></div>
      </div>
    </div>
    <div class="tabs">
      <button class="tab active" onclick="switchTab(this,'all',   '${{name}}')">Tất cả (${{counts.all}})</button>
      <button class="tab"        onclick="switchTab(this,'correct','${{name}}')">✅ Đúng (${{counts.correct}})</button>
      <button class="tab"        onclick="switchTab(this,'wrong',  '${{name}}')">❌ Sai (${{counts.wrong}})</button>
      <button class="tab"        onclick="switchTab(this,'novote', '${{name}}')">⬜ Không vote (${{counts.novote}})</button>
    </div>
    <div id="detail-table-wrap">${{renderTable(rows)}}</div>
  `;
}}

function switchTab(btn, tab, name) {{
  document.querySelectorAll(".tab").forEach(t => t.classList.remove("active"));
  btn.classList.add("active");
  const rows = DATA.details[name] || [];
  const filtered = tab === "all" ? rows : rows.filter(r => r.result === tab);
  document.getElementById("detail-table-wrap").innerHTML = renderTable(filtered);
}}
</script>
</body>
</html>"""

    html_bytes = html.encode("utf-8")
    if output_path:
        with open(output_path, "w", encoding="utf-8") as f:
            f.write(html)
        print(f"Đã xuất: {output_path}")
    return html_bytes


# ─────────────────────────────────────────────────────────────
# Xuất HTML đầy đủ — toàn bộ thành viên hiển thị sẵn
# ─────────────────────────────────────────────────────────────

def export_full_html(
    scores: dict,
    members: list,
    ranking_df: pd.DataFrame,
    member_map: dict,
    output_path: str | None = None,
) -> bytes:
    """
    Tạo HTML báo cáo đầy đủ: bảng xếp hạng + chi tiết từng người.
    Không cần tương tác — cuộn xuống là thấy tất cả.
    """
    generated = __import__("datetime").datetime.now().strftime("%d/%m/%Y %H:%M")
    avg_acc   = ranking_df["Tỷ lệ (%)"].mean()
    top_name  = ranking_df.iloc[0]["Tên"] if not ranking_df.empty else "—"

    # ── Ranking table rows ──
    def medal(rank: int) -> str:
        return {1: "🥇", 2: "🥈", 3: "🥉"}.get(rank, f"#{rank}")

    ranking_rows_html = ""
    for i in range(len(ranking_df)):
        r        = ranking_df.iloc[i]
        rank     = int(ranking_df.index[i])
        bar_w    = int(r["Tỷ lệ (%)"])
        slug     = r["Tên"].replace(" ", "_").replace("'", "")
        ranking_rows_html += f"""
        <tr onclick="document.getElementById('member-{slug}').scrollIntoView({{behavior:'smooth',block:'start'}})"
            style="cursor:pointer">
          <td class="tc fw" style="width:56px">{medal(rank)}</td>
          <td class="fw">{r["Tên"]}</td>
          <td class="tc" style="color:#155724;font-weight:600">{int(r["Đúng"])}</td>
          <td class="tc" style="color:#721c24">{int(r["Sai"])}</td>
          <td style="min-width:140px">
            <div class="bw"><div class="bar" style="width:{bar_w}%"></div></div>
            <span class="bl">{r["Tỷ lệ (%)"]:.1f}%</span>
          </td>
        </tr>"""

    # ── Per-member detail sections ──
    member_sections_html = ""
    for i in range(len(ranking_df)):
        r        = ranking_df.iloc[i]
        rank     = int(ranking_df.index[i])
        name     = r["Tên"]
        member   = member_map.get(name)
        if not member:
            continue

        slug     = name.replace(" ", "_").replace("'", "")
        correct  = int(r["Đúng"])
        wrong    = int(r["Sai"])
        accuracy = r["Tỷ lệ (%)"]

        # build detail rows
        detail_rows = ""
        for match, pred in member["votes"].items():
            if match not in scores:
                continue
            home, away = match.split(" vs ")
            hs, as_    = scores[match]
            result     = grade(pred, match, scores)
            pred_str   = str(pred).strip() if pred else "—"
            if result is True:
                result_html = '<span class="rc">✅ Đúng</span>'
            elif result is False:
                result_html = '<span class="rw">❌ Sai</span>'
            else:
                result_html = '<span class="rn">⬜ Không vote</span>'
            detail_rows += f"""
            <tr>
              <td>{match}</td>
              <td class="tc fw">{hs} - {as_}</td>
              <td>{pred_str}</td>
              <td class="tc">{result_html}</td>
            </tr>"""

        member_sections_html += f"""
      <div class="member-card" id="member-{slug}">
        <div class="member-header">
          <span class="member-rank">{medal(rank)}</span>
          <span class="member-name">{name}</span>
          <div class="member-stats">
            <span class="ms-item ms-correct">✅ {correct} đúng</span>
            <span class="ms-item ms-wrong">❌ {wrong} sai</span>
            <span class="ms-item ms-acc">📊 {accuracy:.1f}%</span>
          </div>
          <button class="toggle-btn" onclick="toggleDetail(this)">▲ Thu gọn</button>
        </div>
        <div class="member-detail">
          <table class="dt">
            <thead><tr>
              <th>Trận</th><th class="tc">Tỉ số</th><th>Dự đoán</th><th class="tc">Kết quả</th>
            </tr></thead>
            <tbody>{detail_rows}</tbody>
          </table>
        </div>
      </div>"""

    html = f"""<!DOCTYPE html>
<html lang="vi">
<head>
<meta charset="UTF-8"/>
<meta name="viewport" content="width=device-width, initial-scale=1.0"/>
<title>⚽ Báo cáo dự đoán bóng đá</title>
<style>
  *, *::before, *::after {{ box-sizing: border-box; margin: 0; padding: 0; }}
  body {{ font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", sans-serif;
          background: #f0f2f6; color: #1a1a2e; }}

  header {{ background: linear-gradient(135deg, #1F4E79, #2e86de); color: #fff;
            padding: 20px 32px; display: flex; align-items: center; gap: 12px;
            position: sticky; top: 0; z-index: 100; box-shadow: 0 2px 12px rgba(0,0,0,.25); }}
  header h1 {{ font-size: 1.4rem; font-weight: 700; }}
  header small {{ margin-left: auto; font-size: .78rem; opacity: .75; }}

  .wrap {{ max-width: 1100px; margin: 0 auto; padding: 20px 16px 48px; }}

  /* Metrics */
  .metrics {{ display: grid; grid-template-columns: repeat(4, 1fr); gap: 12px; margin-bottom: 24px; }}
  .metric  {{ background: #fff; border-radius: 10px; padding: 14px 18px;
              box-shadow: 0 1px 4px rgba(0,0,0,.08); border-left: 4px solid #2e86de; }}
  .metric .ml {{ font-size: .76rem; color: #666; margin-bottom: 4px; }}
  .metric .mv {{ font-size: 1.4rem; font-weight: 700; color: #1F4E79; }}

  /* Section card */
  .section {{ background: #fff; border-radius: 12px;
              box-shadow: 0 1px 4px rgba(0,0,0,.08); margin-bottom: 24px; overflow: hidden; }}
  .section-header {{ background: #1F4E79; color: #fff; padding: 12px 18px;
                     font-weight: 600; font-size: .95rem; display: flex;
                     align-items: center; justify-content: space-between; }}
  .section-header a {{ color: #a8d4ff; font-size: .8rem; text-decoration: none; }}
  .section-header a:hover {{ text-decoration: underline; }}

  /* Ranking table */
  .rt {{ width: 100%; border-collapse: collapse; font-size: .87rem; }}
  .rt thead th {{ background: #EBF3FB; color: #1F4E79; font-weight: 600;
                  padding: 10px 12px; text-align: left; position: sticky; top: 52px; }}
  .rt tbody tr {{ border-bottom: 1px solid #f0f2f6; transition: background .12s; }}
  .rt tbody tr:hover {{ background: #EBF3FB; }}
  .rt td {{ padding: 9px 12px; }}
  .tc {{ text-align: center; }}
  .fw {{ font-weight: 600; }}
  .bw {{ background: #f0f2f6; border-radius: 99px; height: 8px; width: 80px;
         overflow: hidden; display: inline-block; vertical-align: middle; }}
  .bar {{ height: 100%; background: #2e86de; border-radius: 99px; }}
  .bl {{ margin-left: 6px; font-size: .82rem; color: #444; }}

  /* Member cards */
  .members-wrap {{ display: flex; flex-direction: column; gap: 14px; }}
  .member-card {{ background: #fff; border-radius: 12px;
                  box-shadow: 0 1px 4px rgba(0,0,0,.08); overflow: hidden; }}
  .member-header {{ background: linear-gradient(90deg, #1F4E79, #2460a7);
                    color: #fff; padding: 12px 18px; display: flex;
                    align-items: center; gap: 10px; flex-wrap: wrap; cursor: pointer; }}
  .member-rank {{ font-size: 1.2rem; width: 32px; text-align: center; }}
  .member-name {{ font-weight: 700; font-size: 1rem; flex: 1; }}
  .member-stats {{ display: flex; gap: 8px; flex-wrap: wrap; }}
  .ms-item {{ padding: 3px 10px; border-radius: 99px; font-size: .8rem; font-weight: 600; }}
  .ms-correct {{ background: #d4edda; color: #155724; }}
  .ms-wrong   {{ background: #f8d7da; color: #721c24; }}
  .ms-acc     {{ background: rgba(255,255,255,.2); color: #fff; }}
  .toggle-btn {{ margin-left: auto; background: rgba(255,255,255,.15); border: none;
                 color: #fff; padding: 4px 12px; border-radius: 6px; cursor: pointer;
                 font-size: .8rem; white-space: nowrap; }}
  .toggle-btn:hover {{ background: rgba(255,255,255,.28); }}

  .member-detail {{ overflow: hidden; transition: max-height .3s ease; max-height: 2000px; }}
  .member-detail.collapsed {{ max-height: 0; }}

  /* Detail table */
  .dt {{ width: 100%; border-collapse: collapse; font-size: .85rem; }}
  .dt thead th {{ background: #EBF3FB; color: #1F4E79; font-weight: 600;
                  padding: 9px 14px; text-align: left; }}
  .dt tbody tr {{ border-bottom: 1px solid #f0f2f6; }}
  .dt tbody tr:hover {{ background: #fafbfc; }}
  .dt td {{ padding: 8px 14px; }}
  .rc {{ background: #d4edda; color: #155724; font-weight: 600;
         border-radius: 4px; padding: 2px 8px; }}
  .rw {{ background: #f8d7da; color: #721c24; font-weight: 600;
         border-radius: 4px; padding: 2px 8px; }}
  .rn {{ background: #f8f9fa; color: #888; border-radius: 4px; padding: 2px 8px; }}

  .top-link {{ display: inline-block; margin: 8px 0 0 18px; font-size: .78rem;
               color: #2e86de; cursor: pointer; text-decoration: none; }}
  .top-link:hover {{ text-decoration: underline; }}

  @media (max-width: 680px) {{
    .metrics {{ grid-template-columns: 1fr 1fr; }}
    .member-stats {{ display: none; }}
  }}
</style>
</head>
<body id="top">

<header>
  <span style="font-size:1.8rem">⚽</span>
  <h1>Báo cáo dự đoán bóng đá</h1>
  <small>Xuất lúc {generated}</small>
</header>

<div class="wrap">

  <!-- Metrics -->
  <div class="metrics">
    <div class="metric"><div class="ml">Số trận có kết quả</div><div class="mv">{len(scores)}</div></div>
    <div class="metric"><div class="ml">Số người tham gia</div><div class="mv">{len(members)}</div></div>
    <div class="metric"><div class="ml">Tỷ lệ đúng trung bình</div><div class="mv">{avg_acc:.1f}%</div></div>
    <div class="metric"><div class="ml">Người đứng đầu</div><div class="mv" style="font-size:1rem">🥇 {top_name}</div></div>
  </div>

  <!-- Ranking summary -->
  <div class="section">
    <div class="section-header">
      <span>🏆 Bảng xếp hạng</span>
      <a href="#members">Xem chi tiết ↓</a>
    </div>
    <table class="rt">
      <thead><tr>
        <th style="width:56px;text-align:center">Hạng</th>
        <th>Tên</th>
        <th style="text-align:center">Đúng</th>
        <th style="text-align:center">Sai</th>
        <th>Tỷ lệ</th>
      </tr></thead>
      <tbody>{ranking_rows_html}</tbody>
    </table>
  </div>

  <!-- All members detail -->
  <div class="section-header" id="members" style="border-radius:12px 12px 0 0">
    <span>📋 Chi tiết từng người</span>
    <div style="display:flex;gap:8px">
      <button onclick="toggleAll(true)"
              style="background:rgba(255,255,255,.15);border:none;color:#fff;
                     padding:4px 12px;border-radius:6px;cursor:pointer;font-size:.8rem">
        ▲ Thu gọn tất cả
      </button>
      <button onclick="toggleAll(false)"
              style="background:rgba(255,255,255,.15);border:none;color:#fff;
                     padding:4px 12px;border-radius:6px;cursor:pointer;font-size:.8rem">
        ▼ Mở tất cả
      </button>
    </div>
  </div>
  <div class="members-wrap">
    {member_sections_html}
  </div>

  <a href="#top" class="top-link">▲ Lên đầu trang</a>

</div>

<script>
function toggleDetail(btn) {{
  const detail = btn.closest(".member-card").querySelector(".member-detail");
  const collapsed = detail.classList.toggle("collapsed");
  btn.textContent = collapsed ? "▼ Mở rộng" : "▲ Thu gọn";
}}
function toggleAll(collapse) {{
  document.querySelectorAll(".member-detail").forEach(d => {{
    d.classList.toggle("collapsed", collapse);
  }});
  document.querySelectorAll(".toggle-btn").forEach(b => {{
    b.textContent = collapse ? "▼ Mở rộng" : "▲ Thu gọn";
  }});
}}
// Click header để toggle
document.querySelectorAll(".member-header").forEach(h => {{
  h.addEventListener("click", e => {{
    if (!e.target.classList.contains("toggle-btn")) {{
      const btn = h.querySelector(".toggle-btn");
      toggleDetail(btn);
    }}
  }});
}});
</script>
</body>
</html>"""

    html_bytes = html.encode("utf-8")
    if output_path:
        with open(output_path, "w", encoding="utf-8") as f:
            f.write(html)
        print(f"Đã xuất: {output_path}")
    return html_bytes


# ─────────────────────────────────────────────────────────────
# UI
# ─────────────────────────────────────────────────────────────

import subprocess
import sys

# ── Custom CSS ────────────────────────────────────────────────
st.markdown("""
<style>
/* ── Metric cards ── */
[data-testid="metric-container"] {
    background: #f8fafc;
    border: 1px solid #e2e8f0;
    border-radius: 12px;
    padding: 16px 20px;
}
[data-testid="stMetricValue"] { color: #1e40af; font-size: 1.5rem; font-weight: 700; }
[data-testid="stMetricLabel"] { color: #64748b; font-size: .82rem; }

/* ── Tabs ── */
[data-testid="stTabs"] button[role="tab"] { font-weight: 600; }

/* ── Dataframe ── */
[data-testid="stDataFrame"] { border-radius: 10px; overflow: hidden; }

/* ── Containers ── */
[data-testid="stVerticalBlockBorderWrapper"] {
    background: #f8fafc;
    border: 1px solid #e2e8f0 !important;
    border-radius: 12px;
}

/* ── Number input ── */
[data-testid="stNumberInput"] input {
    font-size: 1.2rem;
    font-weight: 700;
    text-align: center;
    border-radius: 8px;
}

/* ── Alert ── */
[data-testid="stAlert"] { border-radius: 10px; }

/* ── Score badge ── */
.score-badge {
    display: block;
    background: #eff6ff;
    border: 2px solid #bfdbfe;
    border-radius: 10px;
    padding: 6px 0;
    font-size: 1.4rem;
    font-weight: 800;
    color: #1e40af;
    letter-spacing: 4px;
    text-align: center;
    margin-top: 8px;
}
.match-title {
    font-size: .88rem;
    font-weight: 600;
    color: #475569;
    text-align: center;
    margin-bottom: 8px;
}
.vs-text {
    font-size: .8rem;
    color: #94a3b8;
    font-weight: 700;
    text-align: center;
    padding-top: 32px;
}
</style>
""", unsafe_allow_html=True)

# ── Sidebar ───────────────────────────────────────────────────
with st.sidebar:
    st.markdown("## ⚽ WC Dashboard")
    st.markdown("---")

    page = st.radio(
        "Chọn trang",
        ["🏆 Bảng xếp hạng", "⚽ Quản lý trận đấu", "📤 Xuất báo cáo"],
        label_visibility="collapsed",
    )

    st.markdown("---")
    st.markdown("### 🔄 Cập nhật dữ liệu")

    # ── Nút 1: Kéo poll Telegram ──
    if st.button("📡 Kéo poll Telegram", use_container_width=True,
                 help="Lấy lại vote từ Telegram → poll_results_1.xlsx"):
        with st.spinner("Đang kéo poll từ Telegram..."):
            result = subprocess.run(
                [sys.executable, "telegram_poll_export.py",
                 "--skip-ranking", "--skip-fetch-scores"],
                capture_output=True, text=True, cwd="."
            )
        output = (result.stdout or "") + (result.stderr or "")
        # Coi là thành công nếu có file được tạo, dù returncode != 0
        # (có thể có PollVoteRequiredError đã được skip trong code)
        import os
        poll_ok = os.path.exists(OUTPUT_FILE)
        if poll_ok:
            st.toast("✅ Kéo poll xong!")
            if output:
                with st.expander("Log", expanded=False):
                    st.code(output, language=None)
            load_all.clear()
            st.rerun()
        else:
            st.error("Kéo poll thất bại")
            st.code(output)

    # ── Nút 2: Fetch kết quả ESPN ──
    if st.button("🌐 Fetch kết quả ESPN", use_container_width=True,
                 help="Tự động lấy tỉ số từ ESPN → scores.json"):
        with st.spinner("Đang fetch từ ESPN..."):
            result = subprocess.run(
                [sys.executable, "telegram_poll_export.py",
                 "--skip-poll", "--fetch-scores", "--skip-ranking"],
                capture_output=True, text=True, cwd="."
            )
        output = (result.stdout or "") + (result.stderr or "")
        if result.returncode == 0:
            st.toast("✅ Fetch kết quả xong!")
        else:
            st.warning("Fetch xong (có thể một số trận chưa khớp)")
        if output:
            with st.expander("Log ESPN", expanded=True):
                st.code(output, language=None)
        load_all.clear()
        st.rerun()

    # ── Nút 3: Cập nhật ranking ──
    if st.button("🏆 Tính lại xếp hạng", use_container_width=True,
                 help="Tính lại BXH từ poll_results_1.xlsx + scores.json hiện tại"):
        load_all.clear()
        st.toast("✅ Đã tính lại xếp hạng!")
        st.rerun()

    st.markdown("---")
    st.caption("Telegram poll & ESPN API")

# ── Load data ─────────────────────────────────────────────────
try:
    match_names, scores, members = load_all()
except FileNotFoundError as e:
    st.error(f"Không tìm thấy file: {e}. Hãy chạy `python telegram_poll_export.py` trước.")
    st.stop()

ranking_df = build_ranking(members, scores)
member_map = {m["name"]: m for m in members}

with open(SCORES_FILE, encoding="utf-8") as f:
    raw_scores = json.load(f)
has_score = {k: v for k, v in raw_scores.items() if v is not None}
no_score  = {k: v for k, v in raw_scores.items() if v is None}

# ── Header metrics ────────────────────────────────────────────
st.markdown("# ⚽ World Cup 2026 — Dự đoán")
st.markdown(f"<p style='color:#64748b;margin-top:-12px;margin-bottom:20px'>Tổng <b style='color:#1e40af'>{len(scores)}</b> trận có kết quả / <b style='color:#d97706'>{len(no_score)}</b> trận chờ</p>", unsafe_allow_html=True)

c1, c2, c3, c4, c5 = st.columns(5)
c1.metric("🏟️ Tổng trận poll",     len(match_names))
c2.metric("✅ Đã có kết quả",       len(scores))
c3.metric("👥 Người tham gia",      len(members))
c4.metric("📊 Tỷ lệ đúng TB",      f"{ranking_df['Tỷ lệ (%)'].mean():.1f}%")
c5.metric("🥇 Dẫn đầu",            ranking_df.iloc[0]["Tên"] if not ranking_df.empty else "—")

st.markdown("<div style='margin-top:8px'></div>", unsafe_allow_html=True)

# ══════════════════════════════════════════════════════════════
# PAGE 1: Bảng xếp hạng
# ══════════════════════════════════════════════════════════════
if page == "🏆 Bảng xếp hạng":

    left, right = st.columns([1, 1.6], gap="large")

    with left:
        st.markdown("### 🏆 Bảng xếp hạng")
        st.caption("Chọn một hàng để xem chi tiết dự đoán →")
        event = st.dataframe(
            ranking_df,
            use_container_width=True,
            selection_mode="single-row",
            on_select="rerun",
            height=640,
            column_config={
                "Tên": st.column_config.TextColumn("Tên", width="medium"),
                "Đúng": st.column_config.NumberColumn("✅ Đúng", width="small"),
                "Sai":  st.column_config.NumberColumn("❌ Sai",  width="small"),
                "Tỷ lệ (%)": st.column_config.NumberColumn("📊 Tỷ lệ", format="%.1f%%"),
            },
        )
        selected_rows = event.selection.rows

    with right:
        if not selected_rows:
            st.markdown("""
            <div style="margin-top:80px;text-align:center;color:#94a3b8">
                <div style="font-size:3rem">👈</div>
                <div style="font-size:1.1rem;margin-top:12px;font-weight:600">Chọn một người trong bảng xếp hạng</div>
                <div style="font-size:.9rem;margin-top:6px">để xem chi tiết từng trận dự đoán</div>
            </div>
            """, unsafe_allow_html=True)
        else:
            row_idx  = selected_rows[0]
            row_data = ranking_df.iloc[row_idx]
            name     = row_data["Tên"]
            rank     = int(ranking_df.index[row_idx])
            member   = member_map[name]

            medal = {1: "🥇", 2: "🥈", 3: "🥉"}.get(rank, f"#{rank}")
            st.markdown(f"### {medal} {name}")

            mc1, mc2, mc3, mc4 = st.columns(4)
            mc1.metric("Hạng",    medal)
            mc2.metric("✅ Đúng", int(row_data["Đúng"]))
            mc3.metric("❌ Sai",  int(row_data["Sai"]))
            mc4.metric("📊 Tỷ lệ", f"{row_data['Tỷ lệ (%)']:.1f}%")

            detail_df = build_detail(member, scores)
            n_correct = (detail_df["Kết quả"] == "✅ Đúng").sum()
            n_wrong   = (detail_df["Kết quả"] == "❌ Sai").sum()
            n_novote  = (detail_df["Kết quả"] == "⬜ Không vote").sum()

            t_all, t_ok, t_bad, t_skip = st.tabs([
                f"Tất cả  {len(detail_df)}",
                f"✅ Đúng  {n_correct}",
                f"❌ Sai  {n_wrong}",
                f"⬜ Bỏ  {n_novote}",
            ])

            def render_detail_cards(df: pd.DataFrame):
                if df.empty:
                    st.info("Không có dữ liệu.")
                    return

                # Header
                st.markdown("""
                <div style="display:grid;grid-template-columns:1fr 80px 120px 110px;gap:12px;
                            padding:6px 16px 8px;color:#94a3b8;font-size:.75rem;font-weight:700;
                            text-transform:uppercase;letter-spacing:.06em;border-bottom:1px solid #e2e8f0;margin-bottom:6px">
                    <span>Trận</span>
                    <span style="text-align:center;display:block">Tỉ số</span>
                    <span style="text-align:center;display:block">Dự đoán</span>
                    <span style="text-align:right;display:block">Kết quả</span>
                </div>
                """, unsafe_allow_html=True)

                # Từng row riêng biệt
                for _, row in df.iterrows():
                    kq = str(row["Kết quả"])
                    if "✅" in kq:
                        badge  = "<span style='background:#dcfce7;color:#166534;padding:3px 12px;border-radius:99px;font-size:.8rem;font-weight:700;white-space:nowrap'>✅ Đúng</span>"
                        border = "#bbf7d0"
                        bg     = "#f0fdf4"
                    elif "❌" in kq:
                        badge  = "<span style='background:#fee2e2;color:#991b1b;padding:3px 12px;border-radius:99px;font-size:.8rem;font-weight:700;white-space:nowrap'>❌ Sai</span>"
                        border = "#fecaca"
                        bg     = "#fff5f5"
                    else:
                        badge  = "<span style='background:#f1f5f9;color:#64748b;padding:3px 12px;border-radius:99px;font-size:.8rem;font-weight:700;white-space:nowrap'>⬜ Bỏ</span>"
                        border = "#e2e8f0"
                        bg     = "#f8fafc"

                    st.markdown(f"""
                    <div style="display:grid;grid-template-columns:1fr 80px 120px 110px;
                                align-items:center;gap:12px;
                                background:{bg};border:1px solid {border};
                                border-radius:10px;padding:10px 16px;margin-bottom:5px">
                        <div style="font-weight:600;font-size:.88rem;color:#1e293b">{row["Trận"]}</div>
                        <div style="text-align:center;font-weight:800;font-size:.95rem;
                                    color:#1e40af;background:#eff6ff;border-radius:6px;padding:3px 6px">{row["Tỉ số"]}</div>
                        <div style="text-align:center;color:#475569;font-size:.85rem">{row["Dự đoán"]}</div>
                        <div style="text-align:right">{badge}</div>
                    </div>
                    """, unsafe_allow_html=True)

            with t_all:   render_detail_cards(detail_df)
            with t_ok:    render_detail_cards(detail_df[detail_df["Kết quả"] == "✅ Đúng"].reset_index(drop=True))
            with t_bad:   render_detail_cards(detail_df[detail_df["Kết quả"] == "❌ Sai"].reset_index(drop=True))
            with t_skip:  render_detail_cards(detail_df[detail_df["Kết quả"] == "⬜ Không vote"].reset_index(drop=True))

# ══════════════════════════════════════════════════════════════
# PAGE 2: Quản lý trận đấu
# ══════════════════════════════════════════════════════════════
elif page == "⚽ Quản lý trận đấu":

    st.markdown("### ⚽ Quản lý tỉ số trận đấu")

    # Progress bar tổng quan
    total_matches = len(raw_scores)
    done          = len(has_score)
    pct           = done / total_matches if total_matches else 0
    st.markdown(f"""
    <div style="margin-bottom:16px">
        <div style="display:flex;justify-content:space-between;color:#64748b;font-size:.85rem;margin-bottom:6px">
            <span>Tiến độ nhập kết quả</span>
            <span><b style="color:#1e40af">{done}</b> / {total_matches} trận</span>
        </div>
        <div style="background:#e2e8f0;border-radius:99px;height:10px;overflow:hidden">
            <div style="background:linear-gradient(90deg,#1e40af,#3b82f6);height:100%;width:{pct*100:.1f}%;border-radius:99px;transition:width .4s"></div>
        </div>
    </div>
    """, unsafe_allow_html=True)

    m_tab1, m_tab2 = st.tabs([
        f"✅ Đã có tỉ số  ({len(has_score)})",
        f"⏳ Chưa có tỉ số  ({len(no_score)})",
    ])

    def render_score_editor(match_dict: dict, key_prefix: str):
        if not match_dict:
            st.markdown("""
            <div style="text-align:center;color:#94a3b8;padding:48px">
                <div style="font-size:2rem">🎉</div>
                <div style="margin-top:8px">Tất cả trận đã có kết quả!</div>
            </div>""", unsafe_allow_html=True)
            return

        matches = list(match_dict.items())
        edits   = {}
        COLS    = 3

        for i in range(0, len(matches), COLS):
            cols = st.columns(COLS)
            for j, col in enumerate(cols):
                if i + j >= len(matches):
                    break
                match, score = matches[i + j]
                home, away   = match.split(" vs ")
                with col:
                    with st.container(border=True):
                        st.markdown(f"<div class='match-title'>{home} <span style='color:#64748b'>vs</span> {away}</div>", unsafe_allow_html=True)
                        ic1, ic2, ic3 = st.columns([5, 2, 5])
                        with ic1:
                            home_val = st.number_input(
                                home, min_value=0, max_value=30,
                                value=int(score[0]) if score else 0,
                                key=f"{key_prefix}_{i+j}_h",
                                label_visibility="collapsed",
                            )
                        with ic2:
                            st.markdown("<div class='vs-text'>—</div>", unsafe_allow_html=True)
                        with ic3:
                            away_val = st.number_input(
                                away, min_value=0, max_value=30,
                                value=int(score[1]) if score else 0,
                                key=f"{key_prefix}_{i+j}_a",
                                label_visibility="collapsed",
                            )
                        st.markdown(
                            f"<div class='score-badge'>{home_val} – {away_val}</div>",
                            unsafe_allow_html=True,
                        )
                        edits[match] = [home_val, away_val]

        st.markdown("<div style='margin-top:16px'></div>", unsafe_allow_html=True)
        sc1, sc2, _ = st.columns([1, 1, 4])
        with sc1:
            if st.button("💾 Lưu thay đổi", key=f"save_{key_prefix}", type="primary", use_container_width=True):
                with open(SCORES_FILE, encoding="utf-8") as f:
                    current = json.load(f)
                current.update(edits)
                with open(SCORES_FILE, "w", encoding="utf-8") as f:
                    json.dump(current, f, ensure_ascii=False, indent=2)
                st.toast(f"✅ Đã lưu {len(edits)} trận!", icon="💾")
                load_all.clear()
                st.rerun()
        with sc2:
            if st.button("↺ Reset", key=f"reset_{key_prefix}", use_container_width=True):
                st.rerun()

    with m_tab1:
        render_score_editor(has_score, "has")
    with m_tab2:
        if no_score:
            st.info(f"⚠️ Còn **{len(no_score)}** trận chưa có kết quả. Nhập tỉ số và nhấn Lưu.")
        render_score_editor(no_score, "no")

# ══════════════════════════════════════════════════════════════
# PAGE 3: Xuất báo cáo
# ══════════════════════════════════════════════════════════════
elif page == "📤 Xuất báo cáo":

    st.markdown("### 📤 Xuất báo cáo HTML")
    st.markdown("<p style='color:#64748b'>File HTML tự chứa — gửi cho bất kỳ ai mà không cần cài đặt gì.</p>", unsafe_allow_html=True)

    ec1, ec2 = st.columns(2)

    with ec1:
        with st.container(border=True):
            st.markdown("#### 🖱️ HTML tương tác")
            st.markdown("<p style='color:#94a3b8;font-size:.88rem'>Bảng xếp hạng + click để xem chi tiết từng người.</p>", unsafe_allow_html=True)
            if st.button("Tạo file", key="gen_interactive", use_container_width=True):
                html_bytes = export_html(match_names, scores, members, ranking_df, member_map)
                st.download_button(
                    label="⬇️ Tải report.html",
                    data=html_bytes,
                    file_name="report.html",
                    mime="text/html",
                    use_container_width=True,
                )

    with ec2:
        with st.container(border=True):
            st.markdown("#### 📄 HTML đầy đủ")
            st.markdown("<p style='color:#94a3b8;font-size:.88rem'>Toàn bộ trang — tất cả chi tiết hiển thị sẵn, cuộn xuống là thấy.</p>", unsafe_allow_html=True)
            if st.button("Tạo file", key="gen_full", use_container_width=True):
                html_bytes = export_full_html(scores, members, ranking_df, member_map)
                st.download_button(
                    label="⬇️ Tải report_full.html",
                    data=html_bytes,
                    file_name="report_full.html",
                    mime="text/html",
                    use_container_width=True,
                )
