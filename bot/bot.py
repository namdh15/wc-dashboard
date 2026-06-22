"""
Telegram bot tổng hợp kết quả poll.
Lệnh: /tonghop — gửi file Excel vào chat hiện tại.
"""
import asyncio
import io
import os
import logging

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import datetime

from telethon import TelegramClient, events
from telethon.sessions import StringSession
from telethon.tl.functions.messages import GetPollVotesRequest
from telethon.tl.types import MessageMediaPoll

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
log = logging.getLogger(__name__)

# ── Cấu hình từ biến môi trường ──────────────────────────────────────────────
API_ID         = int(os.environ["API_ID"])
API_HASH       = os.environ["API_HASH"]
SESSION_STRING = os.environ["SESSION_STRING"]
BOT_TOKEN      = os.environ["BOT_TOKEN"]

# ID group & topic cố định (hoặc cho phép người dùng override sau)
CHAT_ID  = int(os.environ.get("CHAT_ID",  "-1003904575691"))
TOPIC_ID = int(os.environ.get("TOPIC_ID", "3"))
# ─────────────────────────────────────────────────────────────────────────────

HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF")
ROW_FILLS   = [
    PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid"),
    PatternFill(start_color="EBF3FB", end_color="EBF3FB", fill_type="solid"),
]
EMPTY_FILL  = PatternFill(start_color="F5F5F5", end_color="F5F5F5", fill_type="solid")


# ── Telethon helpers ──────────────────────────────────────────────────────────

async def get_voters(client, chat, message, option_bytes, option_text):
    voters, offset = {}, None
    while True:
        res = await client(GetPollVotesRequest(
            peer=chat, id=message.id,
            option=option_bytes, offset=offset, limit=100,
        ))
        if not res.users:
            break
        for u in res.users:
            voters[u.id] = {
                "name":     " ".join(filter(None, [u.first_name, u.last_name])),
                "username": f"@{u.username}" if u.username else "",
                "option":   option_text,
            }
        offset = res.next_offset
        if not offset:
            break
    return voters


async def fetch_poll_data(user_client):
    """Lấy toàn bộ dữ liệu poll từ topic, trả về list poll_data và dict members."""
    chat = await user_client.get_entity(CHAT_ID)

    polls = []
    async for msg in user_client.iter_messages(chat, reply_to=TOPIC_ID, limit=None):
        if msg.media and isinstance(msg.media, MessageMediaPoll):
            polls.append(msg)
    try:
        root = await user_client.get_messages(chat, ids=TOPIC_ID)
        if root and isinstance(root.media, MessageMediaPoll):
            polls.insert(0, root)
    except Exception:
        pass

    polls.sort(key=lambda m: m.id)

    poll_data, members = [], {}
    for msg in polls:
        poll     = msg.media.poll
        results  = msg.media.results
        question = poll.question.text if hasattr(poll.question, "text") else str(poll.question)
        log.info(f"Poll: {question}")

        all_voters = {}
        for answer in poll.answers:
            text = answer.text.text if hasattr(answer.text, "text") else str(answer.text)
            if poll.public_voters:
                v = await get_voters(user_client, chat, msg, answer.option, text)
                all_voters.update(v)

        for uid, info in all_voters.items():
            if uid not in members:
                members[uid] = {"name": info["name"], "username": info["username"]}

        poll_data.append({"question": question, "voters": all_voters, "results": results, "poll": poll})

    return poll_data, members


def build_excel(poll_data, members) -> bytes:
    """Tạo file Excel trong bộ nhớ, trả về bytes."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Kết quả vote"

    num_polls = len(poll_data)

    # Header row 1
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2 + num_polls)
    c = ws.cell(row=1, column=1,
                value=f"Tổng hợp kết quả vote — Xuất lúc: {datetime.now().strftime('%d/%m/%Y %H:%M')}")
    c.font      = Font(bold=True, size=13, color="FFFFFF")
    c.fill      = HEADER_FILL
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26

    # Header row 2
    for col, h in enumerate(["STT", "Họ và tên"], 1):
        c = ws.cell(row=2, column=col, value=h)
        c.font, c.fill = HEADER_FONT, HEADER_FILL
        c.alignment = Alignment(horizontal="center")

    for col_idx, pd in enumerate(poll_data, 3):
        c = ws.cell(row=2, column=col_idx, value=pd["question"])
        c.font      = HEADER_FONT
        c.fill      = HEADER_FILL
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    ws.row_dimensions[2].height = 40

    # Data rows
    member_list = list(members.items())
    for row_idx, (uid, info) in enumerate(member_list, 1):
        row  = row_idx + 2
        fill = ROW_FILLS[row_idx % 2]

        ws.cell(row=row, column=1, value=row_idx).alignment = Alignment(horizontal="center")
        ws.cell(row=row, column=1).fill = fill

        name_display = info["name"]
        if info["username"]:
            name_display += f"\n{info['username']}"
        c = ws.cell(row=row, column=2, value=name_display)
        c.fill      = fill
        c.alignment = Alignment(wrap_text=True)

        for col_idx, pd in enumerate(poll_data, 3):
            voter_info = pd["voters"].get(uid)
            if voter_info:
                c = ws.cell(row=row, column=col_idx, value=voter_info["option"])
                c.fill      = fill
                c.alignment = Alignment(horizontal="center", wrap_text=True)
            else:
                ws.cell(row=row, column=col_idx).fill = EMPTY_FILL

    # Hàng tổng
    total_row = len(member_list) + 3
    ws.cell(row=total_row, column=1, value="Tổng").font      = Font(bold=True)
    ws.cell(row=total_row, column=2, value=f"{len(member_list)} người").font = Font(bold=True)
    for col_idx, pd in enumerate(poll_data, 3):
        c = ws.cell(row=total_row, column=col_idx, value=f"{len(pd['voters'])} phiếu")
        c.font      = Font(bold=True)
        c.alignment = Alignment(horizontal="center")

    # Độ rộng cột
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 28
    for col_idx in range(3, 3 + num_polls):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = 30

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── Bot logic ─────────────────────────────────────────────────────────────────

async def run_bot():
    # Client 1: user session (để gọi Telethon API lấy poll voters)
    user_client = TelegramClient(StringSession(SESSION_STRING), API_ID, API_HASH)
    await user_client.connect()
    if not await user_client.is_user_authorized():
        raise RuntimeError("Session string không hợp lệ. Chạy lại gen_session.py.")

    # Client 2: bot token (để nhận lệnh và gửi file)
    bot = TelegramClient("bot_session", API_ID, API_HASH)
    await bot.start(bot_token=BOT_TOKEN)

    log.info("Bot đã khởi động.")

    @bot.on(events.NewMessage(pattern=r"^/tonghop(@\w+)?$"))
    async def handle_tonghop(event):
        await event.reply("Đang tổng hợp kết quả, vui lòng chờ...")
        try:
            poll_data, members = await fetch_poll_data(user_client)
            if not poll_data:
                await event.reply("Không tìm thấy poll nào trong topic.")
                return
            excel_bytes = build_excel(poll_data, members)
            filename    = f"poll_results_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
            await bot.send_file(
                event.chat_id,
                file=io.BytesIO(excel_bytes),
                attributes=[],
                force_document=True,
                file_name=filename,
                caption=f"Tổng hợp {len(poll_data)} poll — {len(members)} người tham gia",
                reply_to=event.message.id,
            )
        except Exception as e:
            log.exception("Lỗi khi tổng hợp")
            await event.reply(f"Lỗi: {e}")

    await bot.run_until_disconnected()


if __name__ == "__main__":
    asyncio.run(run_bot())
