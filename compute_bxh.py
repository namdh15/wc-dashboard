from openpyxl import load_workbook
import pandas as pd
import re

wb = load_workbook("poll_results_1.xlsx", data_only=True)
ws = wb["Kết quả vote"]
rows = list(ws.iter_rows(values_only=True))
headers = rows[1]

results = {
    "Mexico vs South Africa": (2, 0),
    "South Korea vs Czechia": (2, 1),
    "Canada vs Bosnia": (1, 1),
    "USA vs Paraguay": (4, 1),
    "Qatar vs Switzerland": (1, 1),
    "Brazil vs Morocco": (1, 1),
    "Haiti vs Scotland": (0, 1),
    "Australia vs Türkiye": (2, 0),
    "Germany vs Curacao": (7, 1),
    "Netherlands vs Japan": (2, 2),
    "Ivory Coast vs Ecuador": (1, 0),
    "Sweden vs Tunisia": (5, 1),
    "Spain vs Cape Verde": (0, 0),
    "Belgium vs Egypt": (1, 1),
    "Saudi Arabia vs Uruguay": (1, 1),
    "Iran vs New Zealand": (2, 2),
    "France vs Senegal": (3, 1),
    "Iraq vs Norway": (1, 4),
    "Argentina vs Algeria": (3, 0),
    "Austria vs Jordan": (3, 1),
    "Portugal vs DR Congo": (1, 1),
    "England vs Croatia": (4, 2),
    "Ghana vs Panama": (1, 0),
    # "Uzbekistan vs Columbia":(1, 1),
}

TOTAL = len(results)


def grade(pred, match):
    """
    Trả về:
    1 = dự đoán đúng
    0 = dự đoán sai
    Không vote cũng tính là sai
    """
    if pred is None or str(pred).strip() == "":
        return 0

    pred = str(pred).strip()

    m = re.match(r"(.+?)\s+([+-]\d+(?:\.\d+)?)$", pred)
    if not m:
        return 0

    team = m.group(1).strip()
    handicap = float(m.group(2))

    home, away = match.split(" vs ")
    home_score, away_score = results[match]

    if team == home:
        diff = home_score + handicap - away_score
    elif team == away:
        diff = away_score + handicap - home_score
    else:
        return 0

    return 1 if diff > 0 else 0


scores = []

# Bỏ qua 2 dòng đầu
for row in rows[2:]:
    if not row[1]:
        continue

    name = str(row[1]).split("\n")[0]

    correct = 0

    for col_idx, match in enumerate(headers[2:], start=2):
        if match not in results:
            continue

        pred = row[col_idx]
        correct += grade(pred, match)

    wrong = TOTAL - correct
    accuracy = round(correct / TOTAL * 100, 1)

    scores.append({
        "Tên": name,
        "Đúng": correct,
        "Sai": wrong,
        "Tỷ lệ (%)": accuracy
    })

# Tạo bảng xếp hạng
df = pd.DataFrame(scores)
df = df.sort_values(
    ["Đúng", "Sai", "Tên"],
    ascending=[False, True, True]
).reset_index(drop=True)

df.index += 1
df.index.name = "Hạng"

# In Top 20
print(df.head(20))

# Xuất ra Excel
df.to_excel("ranking.xlsx")
print("\nĐã lưu: ranking.xlsx")